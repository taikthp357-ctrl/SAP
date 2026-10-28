# -*- coding: utf-8 -*-
"""Bộ phân giải mã hàng thông minh từ Sheet DATA tổng hợp (Material Resolver)."""

import os
import re
from pathlib import Path
from typing import List, Dict, Optional, Tuple, Any
import openpyxl


class MaterialResolver:
    """Tự động đối chiếu và chuyển đổi mã rút gọn (vd: 32013765) sang mã SAP đầy đủ (vd: BIW32013765AA)."""

    def __init__(self, excel_path: Optional[Path] = None):
        if excel_path is None:
            excel_path = Path(__file__).resolve().parent.parent / "data" / "DATA_MASTER.xlsx"
        self.excel_path = Path(excel_path)
        self.full_to_meta: Dict[str, Dict[str, Any]] = {}
        self.short_to_fulls: Dict[str, List[str]] = {}
        self.last_mtime: float = 0.0
        self.load_master_data()

    def load_master_data(self):
        """Đọc danh mục mã từ sheet DATA của file Excel."""
        if not self.excel_path.exists():
            print(f"[RESOLVER WARNING] Không tìm thấy file master: {self.excel_path}")
            return

        try:
            current_mtime = os.path.getmtime(self.excel_path)
            if current_mtime == self.last_mtime and self.full_to_meta:
                return  # File chưa thay đổi, không cần nạp lại

            self.full_to_meta.clear()
            self.short_to_fulls.clear()

            wb = openpyxl.load_workbook(str(self.excel_path), read_only=True)
            try:
                sheet_name = "DATA" if "DATA" in wb.sheetnames else wb.sheetnames[0]
                ws = wb[sheet_name]

                # Đọc headers để xác định cột
                rows_iter = ws.iter_rows(values_only=True)
                header_row = next(rows_iter, None)
                if not header_row:
                    return

                full_col_idx = 1
                short_col_idx = 2
                model_col_idx = 3

                for idx, h in enumerate(header_row):
                    if not h:
                        continue
                    h_str = str(h).strip().upper()
                    if "ĐẦY ĐỦ" in h_str or "MATNR" in h_str or h_str == "MÃ":
                        full_col_idx = idx
                    elif "RÚT GỌN" in h_str or "SHORT" in h_str:
                        short_col_idx = idx
                    elif "MODEL" in h_str or "XE" in h_str:
                        model_col_idx = idx

                count = 0
                for row in rows_iter:
                    if not row or len(row) <= full_col_idx:
                        continue
                    full_val = str(row[full_col_idx]).strip().upper() if row[full_col_idx] else ""
                    if not full_val:
                        continue

                    short_val = ""
                    if len(row) > short_col_idx and row[short_col_idx]:
                        short_val = str(row[short_col_idx]).strip().upper()
                    if not short_val:
                        # Tự trích xuất 8 số ở giữa nếu có cấu trúc BIWxxxxxxxxYY
                        if full_val.startswith("BIW") and len(full_val) >= 12:
                            short_val = full_val[3:-2]
                        else:
                            digits = re.findall(r"\d{7,10}", full_val)
                            if digits:
                                short_val = digits[0]
                            else:
                                short_val = full_val

                    model_val = str(row[model_col_idx]).strip() if len(row) > model_col_idx and row[model_col_idx] else ""

                    self.full_to_meta[full_val] = {
                        "full": full_val,
                        "short": short_val,
                        "model": model_val
                    }

                    if short_val not in self.short_to_fulls:
                        self.short_to_fulls[short_val] = []
                    if full_val not in self.short_to_fulls[short_val]:
                        self.short_to_fulls[short_val].append(full_val)

                    count += 1

                self.last_mtime = current_mtime
                print(f"[RESOLVER] Đã nạp {count} mã vật tư từ sheet DATA ({self.excel_path.name})")
            finally:
                wb.close()

        except Exception as e:
            print(f"[RESOLVER ERROR] Lỗi khi nạp dữ liệu DATA: {e}")

    def resolve(self, raw_input: str) -> Tuple[List[str], str]:
        """
        Phân giải mã người dùng nhập thành danh sách mã SAP đầy đủ.
        Trả về: (danh_sách_mã_sap, ghi_chú_nguồn)
        """
        self.load_master_data()  # Kiểm tra tự động reload nếu file Excel vừa được cập nhật

        clean = raw_input.strip().upper()
        if not clean:
            return [], "Rỗng"

        # 0. Tìm kiếm theo mẫu đại diện (Wildcard: * hoặc ?)
        if "*" in clean or "?" in clean:
            import fnmatch
            matched_fulls = []
            for full_code, meta in self.full_to_meta.items():
                short_code = str(meta.get("short", "")).strip().upper()
                if fnmatch.fnmatch(full_code, clean) or (short_code and fnmatch.fnmatch(short_code, clean)):
                    if full_code not in matched_fulls:
                        matched_fulls.append(full_code)
            if matched_fulls:
                note = f"Khớp {len(matched_fulls)} mã theo mẫu wildcard '{clean}'"
                return matched_fulls, note
            return [clean], f"Mẫu wildcard trực tiếp '{clean}' (chưa có trong file Master)"

        # 1. Khớp chính xác với mã đầy đủ đã có trong Master Data
        if clean in self.full_to_meta:
            meta = self.full_to_meta[clean]
            note = f"Khớp mã đầy đủ (Model: {meta.get('model', 'N/A')})"
            return [clean], note

        # 2. Khớp với mã rút gọn (vd: 32013765)
        if clean in self.short_to_fulls:
            full_codes = self.short_to_fulls[clean]
            note = f"Quy đổi từ mã rút gọn {clean} -> {', '.join(full_codes)}"
            return full_codes, note

        # 3. Tìm kiếm theo cụm số / chuỗi con (Fuzzy lookup)
        # Ví dụ người dùng gõ chỉ 6-8 số nằm bên trong mã
        matched_fulls = []
        if len(clean) >= 5:
            for full_code, meta in self.full_to_meta.items():
                if clean in full_code or clean in meta.get("short", ""):
                    matched_fulls.append(full_code)

        if matched_fulls:
            note = f"Tìm thấy {len(matched_fulls)} mã tương ứng từ cụm '{clean}'"
            return matched_fulls, note

        # 4. Nếu là số có 8 chữ số và chưa có trong file Excel, tự động suy đoán tiền tố BIW...AA
        # 4. Tự động chuẩn hóa tiền tố xưởng (BIW / ASY / TRM) và 8 chữ số gốc
        m_shop = re.fullmatch(r"([A-Za-z]{3})[-_\s]*(\d{8})([A-Za-z]{2})?", clean)
        if m_shop:
            pfx = m_shop.group(1).upper()
            core = m_shop.group(2)
            sfx = m_shop.group(3).upper() if m_shop.group(3) else "AA"
            inferred = f"{pfx}{core}{sfx}"
            wildcard = f"{pfx}{core}*"
            note = f"Tự động chuẩn hóa tiền tố {pfx}: {clean} -> {inferred} / core: {core}"
            return [inferred, core, wildcard], note

        if re.fullmatch(r"\d{8}", clean):
            inferred = f"BIW{clean}AA"
            note = f"Tự động chuẩn hóa: {clean} -> {inferred} (chưa có trong file Master)"
            return [inferred, clean, f"BIW{clean}*"], note

        # 5. Fallback: Giữ nguyên mã người dùng nhập để gửi thẳng vào SAP
        return [clean], "Mã nhập trực tiếp (không qua từ điển)"

    def get_materials_by_model(self, model_name: str, limit: int = 0) -> List[str]:
        """Lấy danh sách mã SAP đầy đủ theo Model xe (vd: Limo, VF5, VF8...)."""
        self.load_master_data()
        clean_model = model_name.strip().upper()
        results = []
        for full_code, meta in self.full_to_meta.items():
            item_model = str(meta.get("model", "")).strip().upper()
            if clean_model in item_model or (clean_model == "LIMO" and "LIMO" in item_model):
                if full_code not in results:
                    results.append(full_code)
                    if limit > 0 and len(results) >= limit:
                        break
        return results

    @staticmethod
    def detect_model_and_shop(text: str) -> Tuple[Optional[str], str, List[str]]:
        """Phân tích văn bản để nhận diện Model và xác định Shop (BODY vs XMĐ vs EDC) cùng danh sách Plant chuẩn.
        
        Trả về: (model_name, shop_name, list_of_plants)
        - EDC:  Plant 4007 & 3093
        - BODY: Plant 3093 & 3001
        - XMĐ:  Plant 3093, 2001 & 3001
        """
        norm = text.upper()
        # 0. Kiểm tra Xưởng Sơn EDC / Plant 4007
        edc_patterns = [
            ("EDC", ["EDC", "SON EDC", "SƠN EDC", "XUONG SON", "XƯỞNG SƠN", "4007", "4004"])
        ]
        for m_name, kws in edc_patterns:
            if any(re.search(rf"\b{re.escape(kw)}\b", norm) for kw in kws):
                return m_name, "EDC", ["4007", "3093"]

        # 1. Kiểm tra XMĐ trước nếu có từ khóa xe máy điện / trạm mạ / xmd
        xmd_patterns = [
            ("XMD", ["XMD", "XE MAY DIEN", "XE MÁY ĐIỆN", "TRAM MA", "TRẠM MẠ", "MA SAU SON", "MÃ SAU SƠN"]),
            ("KLARA", ["KLARA"]),
            ("FELIZ", ["FELIZ"]),
            ("VENTO", ["VENTO"]),
            ("THEON", ["THEON"]),
            ("EVO", ["EVO"])
        ]
        for m_name, kws in xmd_patterns:
            if any(re.search(rf"\b{re.escape(kw)}\b", norm) for kw in kws):
                return m_name, "XMD", ["3093", "2001", "3001"]

        # 2. Kiểm tra các model BODY
        body_patterns = [
            ("LIMO", ["LIMO", "VF8 LIMO", "LIMOUSINE"]),
            ("VF3", ["VF3", "VF 3"]),
            ("VF5", ["VF5", "VF 5", "VF5 RHD", "VF5 LHD"]),
            ("VF6", ["VF6", "VF 6", "VF6 CIS", "VF6 RHD"]),
            ("VF7", ["VF7", "VF 7", "VF7 RHD"]),
            ("VF8", ["VF8", "VF 8", "VF8N", "VF8NP", "VF8 EU"]),
            ("VF9", ["VF9", "VF 9", "VF9NP", "VF9 LH"]),
            ("VFE34", ["VFE34", "VF E34", "E34"]),
            ("VF6,7", ["VF6,7", "VF67"]),
            ("BODY", ["BODY", "THAN VO", "THÂN VỎ", "HAN", "HÀN"])
        ]
        for m_name, kws in body_patterns:
            if any(re.search(rf"\b{re.escape(kw)}\b", norm) for kw in kws):
                return m_name, "BODY", ["3093", "3001"]

        return None, "BODY", ["3093", "3001"]

    def resolve_list(self, raw_inputs: List[str]) -> Tuple[List[str], List[str]]:
        """Phân giải một danh sách mã hàng, loại bỏ trùng lặp."""
        all_resolved = []
        notes = []
        seen = set()

        for item in raw_inputs:
            resolved_codes, note = self.resolve(item)
            notes.append(note)
            for c in resolved_codes:
                if c not in seen:
                    seen.add(c)
                    all_resolved.append(c)

        return all_resolved, notes
