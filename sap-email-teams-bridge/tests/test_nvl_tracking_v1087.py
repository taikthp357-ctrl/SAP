# -*- coding: utf-8 -*-
"""Unit tests cho bộ hoàn thiện các nhánh nghiệp vụ NVL (v1.0.87):
1. Báo cáo tình trạng cấp hàng theo Model/Ca (NvlTrackingService.get_shift_supply_status).
2. Tự động ghi nhận số OB và đồng bộ trạng thái vào file theo dõi (sync_dispatch_to_tracking_file).
3. Hướng dẫn & T-Code in phiếu WHT (/SCWM/MON, /SCWM/PRDO, MB90, LT31).
4. Bóc tách email order sản xuất (parse_production_order_text).
5. Điều phối Router & NvlHandler cho TINH_TRANG_NVL và IN_WHT.
"""

import os
import shutil
import tempfile
from pathlib import Path
from unittest.mock import MagicMock

import pytest
import openpyxl

from app.email_parser import SmartCommandParser
from app.nvl_tracking import NvlTrackingService
from app.handlers.nvl_handler import NvlHandler
from app.router import router
from app.core.context import RequestContext, CommandResult


@pytest.fixture
def mock_tracking_workbook(tmp_path):
    """Tạo file Excel theo dõi mẫu chuẩn cấu hình 37 cột."""
    file_path = tmp_path / "261028_THEO DOI ODER NVL_P_S6.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ODER NVl"

    headers = [
        "NCC", "KV", "Packing", "Mã NVL", "SL oder", "Trạm", "Plant", "Sloc Kho", "Sloc Line", "Model",
        "Ca Nhập", "Ca chạy", "OB", "OB1", "SP_1", "SL_1", "OB2", "SP_2", "SL_2", "OB_3",
        "SP_3", "SL_3", "OB_4", "SP_4", "SL_4", "Total cấp", "Gap", "Tình Trạng cấp đơn", "TT nhặt Hàng", "KHSX",
        "NOTE", "Ngày order", "Giờ order", "Người order", "Ca", "SL San Ca", "Ghi Chú San Ca"
    ]
    ws.append(headers)

    # Dòng 1: VF6 - Ca S6 - Pending (Chưa cấp)
    ws.append([
        "SIMWON", "KV01", 10, "BIW70004393AB", 400, "SUB_SW", 3093, 1005, 3001, "VF6",
        "S6", "S6", "", "", "", 0, "", "", 0, "", "", 0, "", "", 0,
        0, -400, "Pending", "Chưa Cấp", "2026-10-28", "", "2026-10-28", "08:00", "TL_A", "S6", 0, ""
    ])

    # Dòng 2: VF6 - Ca S6 - Partial finished (Cấp thiếu)
    ws.append([
        "0x2a", "KV02", 20, "BIW35113211AB", 80, "MAIN_BODY", 3093, 1005, 3001, "VF6",
        "S6", "S6", "1012750084", "1012750084", "16-01", 50, "", "", 0, "", "", 0, "", "", 0,
        50, -30, "Partial finished", "Cấp thiếu", "2026-10-28", "", "2026-10-28", "08:00", "TL_A", "S6", 0, ""
    ])

    # Dòng 3: LIMO - Ca S2 - Finished (Cấp đủ)
    ws.append([
        "VINFAST", "KV03", 50, "BIW99999999AA", 100, "DOOR_LINE", 3093, 1005, 3001, "LIMO",
        "S6", "S2", "1012750099", "1012750099", "16-02", 100, "", "", 0, "", "", 0, "", "", 0,
        100, 0, "Finished", "Đã cấp", "2026-10-28", "", "2026-10-28", "08:00", "TL_B", "S2", 0, ""
    ])

    wb.save(str(file_path))
    return file_path


def test_parser_tinh_trang_nvl():
    """Kiểm tra SmartCommandParser bóc tách đúng câu lệnh hỏi tình trạng NVL."""
    p1 = SmartCommandParser.parse("tình trạng cấp hàng VF6 ca S6")
    assert p1["command"] == "TINH_TRANG_NVL"
    assert p1["model"] == "VF6"
    assert p1["shift"] == "S6"
    assert p1["valid"] is True

    p2 = SmartCommandParser.parse("báo cáo cấp nvl ca S2 ngày 2026-10-28")
    assert p2["command"] == "TINH_TRANG_NVL"
    assert p2["shift"] == "S2"
    assert "2026-10-28" in p2.get("date", "")
    assert p2["valid"] is True


def test_parser_in_wht():
    """Kiểm tra SmartCommandParser bóc tách lệnh in WHT."""
    p = SmartCommandParser.parse("in wht 1012816323")
    assert p["command"] == "IN_WHT"
    assert p["obNo"] == "1012816323"
    assert p["valid"] is True

    p_guide = SmartCommandParser.parse("tcode in wht")
    assert p_guide["command"] == "IN_WHT"
    assert p_guide["valid"] is True


def test_get_shift_supply_status(mock_tracking_workbook):
    """Kiểm tra báo cáo tiến độ và KPI từ file theo dõi mẫu."""
    status = NvlTrackingService.get_shift_supply_status(
        tracking_file=mock_tracking_workbook,
        shift="S6",
        model="VF6"
    )

    assert status["status"] == "OK"
    kpi = status["kpi"]
    assert kpi["total_orders"] == 2
    assert kpi["completed_count"] == 0
    assert kpi["partial_count"] == 1
    assert kpi["pending_count"] == 1
    assert len(status["shortages"]) == 2

    # Kiểm tra tự động nạp NCC từ Master Data cho ô '0x2a'
    sh1 = next(s for s in status["shortages"] if s["material"] == "BIW35113211AB")
    assert sh1["ncc"] != "0x2a"


def test_sync_dispatch_to_tracking_file(mock_tracking_workbook):
    """Kiểm tra tự động ghi nhận số OB, tính lại Gap và cập nhật trạng thái."""
    test_orders = [
        {"material": "BIW70004393AB", "qty": 400.0, "allocated_qty": 400.0, "station": "SUB_SW"},
        {"material": "BIW35113211AB", "qty": 30.0, "allocated_qty": 30.0, "station": "MAIN_BODY"}
    ]

    res = NvlTrackingService.sync_dispatch_to_tracking_file(
        orders_dispatched=test_orders,
        delivery_no="5000123456",
        shift="S6",
        tracking_file=mock_tracking_workbook
    )

    assert res["status"] == "SUCCEEDED"
    assert res["updated_count"] == 2

    # Đọc lại file sau khi đồng bộ
    wb = openpyxl.load_workbook(str(mock_tracking_workbook), data_only=True)
    ws = wb["ODER NVl"]

    # Dòng 2 (BIW70004393AB): ban đầu Pending -> giờ đủ 400 -> Finished
    assert str(ws.cell(2, 14).value) == "5000123456"
    assert float(ws.cell(2, 16).value) == 400.0
    assert float(ws.cell(2, 26).value) == 400.0
    assert float(ws.cell(2, 27).value) == 0.0
    assert ws.cell(2, 28).value == "Finished"
    assert ws.cell(2, 29).value == "Đã cấp"

    # Dòng 3 (BIW35113211AB): ban đầu có OB1 (50 PC) -> ghi vào OB2 (30 PC) -> Total = 80 -> Finished
    assert str(ws.cell(3, 14).value) == "1012750084"
    assert str(ws.cell(3, 17).value) == "5000123456"
    assert float(ws.cell(3, 19).value) == 30.0
    assert float(ws.cell(3, 26).value) == 80.0
    assert float(ws.cell(3, 27).value) == 0.0
    assert ws.cell(3, 28).value == "Finished"


def test_wht_print_guidance():
    """Kiểm tra chỉ dẫn T-Code in phiếu WHT."""
    res_ewm = NvlTrackingService.get_wht_print_guidance("1012816323")
    assert res_ewm["status"] == "OK"
    assert "/SCWM/MON" in res_ewm["tcodes"]
    assert "/SCWM/PRDO" in res_ewm["tcodes"]

    res_ecc = NvlTrackingService.get_wht_print_guidance("5000123456", system_type="ECC")
    assert "MB90" in res_ecc["tcodes"]
    assert "LT31" in res_ecc["tcodes"]


def test_parse_production_order_text():
    """Kiểm tra bóc tách văn bản order email từ xưởng."""
    mail_body = """
    Gửi kho, chuyền VF6 xin cấp bổ sung:
    BIW70004393AB    400    SUB_SW
    BIW35113211AB    80     MAIN_BODY
    """
    orders = NvlTrackingService.parse_production_order_text(mail_body)
    assert len(orders) == 2
    assert orders[0]["material"] == "BIW70004393AB"
    assert orders[0]["qty"] == 400.0
    assert orders[1]["material"] == "BIW35113211AB"
    assert orders[1]["qty"] == 80.0


def test_nvl_handler_tinh_trang_dispatch(mock_tracking_workbook):
    """Kiểm tra điều phối CommandRouter và NvlHandler xử lý lệnh TINH_TRANG_NVL."""
    req = {
        "command": "TINH_TRANG_NVL",
        "shift": "S6",
        "model": "VF6",
        "tracking_files_dir": str(mock_tracking_workbook.parent)
    }

    ctx = RequestContext(
        request_id="REQ_TEST_TRACKING",
        command="TINH_TRANG_NVL",
        raw_materials=[],
        resolved_materials=[],
        meta=req,
        plant="3093",
        config={"tracking_files_dir": str(mock_tracking_workbook.parent)}
    )

    handler = NvlHandler()
    result = handler.handle(ctx)

    assert result.status == "SUCCEEDED"
    assert "Báo cáo tiến độ cấp NVL" in result.summary
    assert len(result.records) == 2
    assert result.custom_html is not None
    assert "BÁO CÁO TIẾN ĐỘ CẤP HÀNG NVL" in result.custom_html
