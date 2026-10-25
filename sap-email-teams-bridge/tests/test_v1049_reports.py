"""Business regressions: all fixtures are synthetic, exports stay in tmp_path."""
from pathlib import Path
import openpyxl
import pytest

from app.bbbg_service import BbbgService as B
from app.extended_query_service import ExtendedQueryService as Q


def mb(**changes):
    row = dict(product="M1", batch="B1", unit="PC", refNo="18001", qty=10,
               docNo="50001", docYear="2026", docItem="1", mvtType="101")
    row.update(changes)
    return row


def ib(**changes):
    row = dict(product="M1", batch="B1", unit="PC", ibNo="18001", actualQty=10)
    row.update(changes)
    return row


@pytest.mark.parametrize("changes", [dict(ibNo="18002"), dict(product="M2"),
                                      dict(batch="B2"), dict(unit="KG"),
                                      dict(plant="3093"), dict(sloc="1005")])
def test_wrong_key_never_matches(changes):
    result = B.reconcile([mb()], [ib(**changes)])
    assert result["matched_count"] == 0
    assert all(not x["gr_confirmed"] for x in result["items"])


def test_equal_quantity_is_not_gr_evidence():
    result = B.reconcile([mb()], [ib()])
    assert result["matched_count"] == 1
    assert result["requires_review"] is True
    assert result["items"][0]["gr_status"] == "Chưa xác nhận GR"


@pytest.mark.parametrize("qty", [-10, 10])
def test_reversal_reduces_net_quantity(qty):
    result = B.reconcile([mb(), mb(docNo="50002", qty=qty, mvtType="102")], [ib()])
    assert result["matched_count"] == 0
    item = result["items"][0]
    assert item["handover_qty"] == 0
    assert item["diff_qty"] == -10


def test_many_postings_and_one_receipt_aggregate():
    result = B.reconcile([mb(qty=4), mb(qty=6, docNo="50002")], [ib()])
    assert result["matched_count"] == 1
    assert result["items"][0]["source_row_count"]["MB51"] == 2


def test_many_to_many_items_group_without_consuming_other_items():
    result = B.reconcile(
        [mb(qty=4, ibItem="10"), mb(qty=6, docNo="50002", ibItem="10"),
         mb(qty=3, docNo="50003", ibItem="20")],
        [ib(item="10", actualQty=10), ib(item="20", actualQty=3)])
    assert result["matched_count"] == 2
    assert result["total_handover_qty"] == 13


def test_duplicate_document_year_item_is_not_double_counted():
    row = mb()
    result = B.reconcile([row, dict(row)], [ib()])
    assert result["duplicate_count"] == 1
    assert result["total_handover_qty"] == 10
    assert result["matched_count"] == 1


def test_conflicting_same_document_requires_review():
    result = B.reconcile([mb(), mb(qty=20)], [ib()])
    assert result["matched_count"] == 0
    assert result["review_required_count"] > 0


def test_same_document_in_different_year_is_distinct():
    result = B.reconcile([mb(qty=4), mb(qty=6, docYear="2025")], [ib()])
    assert result["matched_count"] == 1
    assert result["duplicate_count"] == 0


def test_missing_item_identity_does_not_silently_sum_duplicates():
    result = B.reconcile([mb(docItem=""), mb(docItem="")], [ib()])
    assert result["matched_count"] == 0
    assert result["review_required_count"] > 0


@pytest.mark.parametrize("qty", [None, "", "bogus", "abc.123"])
def test_invalid_quantity_is_review_not_zero(qty):
    result = B.reconcile([mb(qty=qty)], [ib(actualQty=0)])
    assert result["matched_count"] == 0
    assert result["review_required_count"] > 0


def test_zero_actual_quantity_does_not_fall_back_to_positive_qty():
    result = B.reconcile([mb()], [ib(actualQty=0, qty=10)])
    assert result["matched_count"] == 0
    assert result["items"][0]["received_qty"] == 0


@pytest.mark.parametrize("mvt", ["311", "601"])
def test_non_receipt_movements_never_prove_gr(mvt):
    result = B.reconcile([mb(mvtType=mvt)], [ib()])
    assert result["matched_count"] == 0
    assert result["review_required_count"] > 0


def test_mixed_uom_has_no_meaningless_total():
    result = B.reconcile([mb(), mb(product="M2", unit="KG", docNo="50002")],
                         [ib(), ib(product="M2", unit="KG")])
    assert result["total_handover_qty"] is None
    assert set(result["totals_by_unit"]) == {"PC", "KG"}


def test_empty_export_has_no_circular_total_or_fake_approval(tmp_path):
    path = B.export_excel(B.reconcile([], []), output_dir=tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["BBBG_NEW"]
    assert ws["G13"].value == 0
    assert ws["H13"].value == 0
    assert "DRAFT" in ws["G2"].value
    assert ws["C5"].value is None
    assert ws["G5"].value is None
    assert ws["C9"].value is None
    wb.close()


def test_template_is_explicitly_unsupported_not_ignored(tmp_path):
    with pytest.raises(ValueError, match="template"):
        B.export_excel(B.reconcile([], []), template_path=tmp_path / "template.xlsx", output_dir=tmp_path)


def test_exports_unique_and_batch_cannot_escape_output_dir(tmp_path):
    result = B.reconcile([], [])
    result["batch_id"] = "../../outside"
    paths = [Path(B.export_excel(result, output_dir=tmp_path)) for _ in range(2)]
    assert paths[0] != paths[1]
    assert all(p.parent == tmp_path for p in paths)


def test_manual_excel_text_cannot_be_formula(tmp_path):
    result = B.reconcile([], [], [dict(material="=1+1", qty=1, unit="PC", note="=HYPERLINK(\"bad\")")])
    wb = openpyxl.load_workbook(B.export_excel(result, output_dir=tmp_path), data_only=False)
    assert wb["BBBG_NEW"]["B13"].data_type == "s"
    assert wb["BBBG_NEW"]["K13"].data_type == "s"
    wb.close()


def test_preview_escapes_text_and_does_not_invent_send_workflow():
    result = B.reconcile([], [], [dict(material="<script>bad</script>", qty=1)])
    preview = B.generate_mail_preview(result)
    assert "<script>" not in preview["html_preview"]
    assert "&lt;script&gt;" in preview["html_preview"]
    assert "XÁC NHẬN GỬI" not in preview["html_preview"]
    assert preview["recipients"] == []
    assert preview["send_supported"] is False


def test_asn_filters_correct_inbound_and_does_not_invent_clearance():
    result = Q.track_asn_provenance("18001", [ib(qty=10), ib(ibNo="18002", qty=999, customs_status="CLEARED")])
    assert result["found"] is True
    assert result["total_qty"] == 10
    assert result["customs_status"] == "UNKNOWN"
    assert result["items"][0]["status"] == "UNKNOWN"
    assert result["excluded_rows"] == 1


def test_asn_unidentified_rows_are_not_assigned_requested_inbound():
    result = Q.track_asn_provenance("18001", [dict(product="M1", qty=10)])
    assert result["found"] is False
    for key in ("container_no", "bill_of_lading", "vendor", "eta", "customs_status"):
        assert result[key] == "UNKNOWN"
    assert result["items"] == []


def test_asn_conflicting_header_fields_not_first_row_wins():
    result = Q.track_asn_provenance("18001", [ib(container="C1"), ib(container="C2")])
    assert result["container_no"] == "MULTIPLE_VALUES"
    assert result["field_values"]["container_no"] == ["C1", "C2"]


def test_asn_mixed_units_no_combined_total():
    result = Q.track_asn_provenance("18001", [ib(qty=10), ib(qty=20, unit="KG")])
    assert result["total_qty"] is None
    assert result["totals_by_unit"] == {"PC": 10, "KG": 20}


def test_source_errors_do_not_mean_no_bom_or_queue_clean():
    assert Q.explode_bom_structure("M1", "P1", [], source_complete=False)["has_bom"] is None
    assert Q.diagnose_ewm_queues([], source_error="timeout")["is_clean"] is None
