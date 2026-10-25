# -*- coding: utf-8 -*-
"""Unit tests for BbbgService (Đợt 3: BBBG FG Reconciliation & Template Export)."""

import os
from pathlib import Path
import openpyxl
import pytest

from app.bbbg_service import BbbgService


def test_bbbg_reconcile_matched_and_mismatched():
    mb51_rows = [
        {"product": "BIW32013765AA", "batch": "B2610A", "docNo": "5000123456", "qty": "100", "unit": "PC", "refNo": "1800554433"},
        {"product": "BIW32024300AB", "batch": "B2610B", "docNo": "5000123457", "qty": "50", "unit": "PC", "refNo": "1800554434"},
        {"product": "BIW32030000AA", "batch": "B2610C", "docNo": "5000123458", "qty": "30", "unit": "PC", "refNo": "1800554435"},
    ]

    vl06i_rows = [
        {"product": "BIW32013765AA", "batch": "B2610A", "ibNo": "1800554433", "actualQty": "100", "unit": "PC"},  # MATCHED
        {"product": "BIW32024300AB", "batch": "B2610B", "ibNo": "1800554434", "actualQty": "45", "unit": "PC"},   # MISMATCH (50 vs 45)
        # BIW32030000AA missing in VL06I -> MISSING_INBOUND
        {"product": "BIW99999999XX", "batch": "B9999X", "ibNo": "1800554999", "actualQty": "20", "unit": "PC"},   # Missing in MB51 -> MISSING_MB51
    ]

    manual_rows = [
        {"material": "NON_SAP_TOOL_01", "batch": "M01", "qty": 5, "note": "Gá mẫu thử nghiệm ngoài hệ thống"}
    ]

    result = BbbgService.reconcile(mb51_rows, vl06i_rows, manual_rows)

    assert result["total_items"] == 5
    assert result["matched_count"] == 1
    assert result["mismatch_count"] == 1
    assert result["missing_inbound_count"] == 1
    assert result["missing_mb51_count"] == 1
    assert result["manual_count"] == 1
    assert result["has_mismatch"] is True

    # Check matched item
    matched = [it for it in result["items"] if it["status"] == "MATCHED"][0]
    assert matched["material"] == "BIW32013765AA"
    assert matched["diff_qty"] == 0.0

    # Check mismatch item
    mismatch = [it for it in result["items"] if it["status"] == "MISMATCH_QTY"][0]
    assert mismatch["material"] == "BIW32024300AB"
    assert mismatch["diff_qty"] == 5.0

    # Check manual item
    manual = [it for it in result["items"] if it["status"] == "MANUAL_ENTRY"][0]
    assert manual["is_manual"] is True
    assert manual["material"] == "NON_SAP_TOOL_01"


def test_bbbg_export_excel(tmp_path):
    mb51_rows = [
        {"product": "BIW32013765AA", "batch": "B2610A", "docNo": "5000123456", "qty": "100", "unit": "PC", "refNo": "1800554433"},
        {"product": "BIW32024300AB", "batch": "B2610B", "docNo": "5000123457", "qty": "50", "unit": "PC", "refNo": "1800554434"},
    ]
    vl06i_rows = [
        {"product": "BIW32013765AA", "batch": "B2610A", "ibNo": "1800554433", "actualQty": "100", "unit": "PC"},
        {"product": "BIW32024300AB", "batch": "B2610B", "ibNo": "1800554434", "actualQty": "50", "unit": "PC"},
    ]

    result = BbbgService.reconcile(mb51_rows, vl06i_rows)
    export_path = BbbgService.export_excel(result, output_dir=tmp_path)

    assert os.path.exists(export_path)
    wb = openpyxl.load_workbook(export_path, data_only=False)
    assert "BBBG_NEW" in wb.sheetnames
    ws = wb["BBBG_NEW"]

    # Header check
    assert "BIÊN BẢN GIAO NHẬN" in str(ws["C2"].value)
    assert ws["A12"].value == "TT\n(No)"
    assert ws["B12"].value == "Mã Hàng Hóa\n(Part No.)"

    # Data check
    assert ws.cell(row=13, column=2).value == "BIW32013765AA"
    assert ws.cell(row=13, column=7).value == 100.0
    assert ws.cell(row=13, column=8).value == 100.0

    # Total check
    assert ws.cell(row=15, column=1).value == "TỔNG CỘNG"
    assert "=SUM(G13:G14)" in str(ws.cell(row=15, column=7).value)
    assert "=SUM(H13:H14)" in str(ws.cell(row=15, column=8).value)


def test_bbbg_mail_preview():
    mb51_rows = [
        {"product": "BIW32013765AA", "batch": "B2610A", "docNo": "5000123456", "qty": "100", "unit": "PC", "refNo": "1800554433"}
    ]
    vl06i_rows = [
        {"product": "BIW32013765AA", "batch": "B2610A", "ibNo": "1800554433", "actualQty": "100", "unit": "PC"}
    ]
    result = BbbgService.reconcile(mb51_rows, vl06i_rows)
    preview = BbbgService.generate_mail_preview(result, shop_key="BODY")

    assert preview["requires_confirmation"] is True
    assert preview["shop"] == "Xưởng Hàn (Body)"
    assert preview["recipients"] == []
    assert preview["send_supported"] is False
    assert "XEM TRƯỚC BIÊN BẢN BÀN GIAO" in preview["html_preview"]
    assert "BIW32013765AA" in preview["html_preview"]
