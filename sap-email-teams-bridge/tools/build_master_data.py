# -*- coding: utf-8 -*-
"""Tool to generate unified MASTER_DATA_VINFAST.xlsx and sync to data/nvl_master_v31.json.

Combines:
- 260810_RA_DON_NVL_V31_ZONE.xlsm (Catalog, Box packing, Bin patterns, Config)
- THEO DOI ODER NVL_P_*.xlsx (NCC vendor mappings for 850+ materials)
- 261008_Stock_FG_MA_BODY.xlsm (FG Catalog, Model, US usage, Trạm, Shop)
- VinFast Warehouse Physical Bin & FEFO Rules (P22/23 RV5, P24 Orange Rack, P25, P26, ZF, HR1, SLoc 200C)
"""

import glob
import json
import os
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def safe_int_packing(val) -> int:
    if not val:
        return 0
    s = str(val).strip()
    try:
        return int(float(s))
    except Exception:
        return 0


def build_master_data():
    base_dir = Path("G:/SAP")
    bridge_dir = base_dir / "sap-email-teams-bridge"
    output_excel = base_dir / "MASTER_DATA_VINFAST.xlsx"
    output_json = bridge_dir / "data" / "nvl_master_v31.json"

    print(f"[BUILD_MASTER] Starting aggregation...")

    # 1. Extract NCC and metadata from all THEO DOI ODER NVL files
    ncc_files = sorted(glob.glob(str(base_dir / "*THEO DOI ODER NVL*.xlsx")))
    ncc_map = {}
    material_meta_from_td = {}
    fefo_extra_from_td = set()

    print(f"[BUILD_MASTER] Scanning {len(ncc_files)} THEO DOI ODER NVL files...")
    for fpath in ncc_files:
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            ws = wb.active
            for r in range(2, ws.max_row + 1):
                try:
                    ncc = ws.cell(r, 1).value
                    mat = ws.cell(r, 4).value
                    if not mat:
                        continue
                    m_clean = str(mat).strip().upper()
                    if not m_clean:
                        continue

                    pkg_val = ws.cell(r, 3).value
                    pkg_int = safe_int_packing(pkg_val)
                    if pkg_val and "HSD" in str(pkg_val).upper():
                        fefo_extra_from_td.add(m_clean)

                    if m_clean not in material_meta_from_td:
                        kv = str(ws.cell(r, 2).value or "").strip()
                        tram = str(ws.cell(r, 6).value or "").strip()
                        model = str(ws.cell(r, 10).value or "").strip()
                        material_meta_from_td[m_clean] = {
                            "kv": kv,
                            "packing": pkg_int,
                            "tram": tram,
                            "model": model,
                        }
                    elif pkg_int > 0 and material_meta_from_td[m_clean]["packing"] == 0:
                        material_meta_from_td[m_clean]["packing"] = pkg_int

                    if ncc and str(ncc).strip() not in ("0x2a", "None", "", "nan"):
                        ncc_str = str(ncc).strip()
                        if m_clean not in ncc_map:
                            ncc_map[m_clean] = ncc_str
                except Exception:
                    pass
        except Exception as e:
            print(f"  Warning reading {fpath}: {e}")

    print(f"[BUILD_MASTER] Found {len(ncc_map)} material-to-vendor mappings.")

    # 2. Extract from 260810_RA_DON_NVL_V31_ZONE.xlsm
    v31_path = base_dir / "260810_RA_DON_NVL_V31_ZONE.xlsm"
    wb_v31 = openpyxl.load_workbook(str(v31_path), data_only=True)
    ws_dm = wb_v31["DANH_MUC"]
    ws_cfg = wb_v31["CONFIG"]

    # FEFO list
    fefo_materials = {}
    for r in range(5, ws_dm.max_row + 1):
        p = ws_dm.cell(r, 1).value
        t = ws_dm.cell(r, 2).value
        if p:
            fefo_materials[str(p).strip().upper()] = str(t or "HSD").strip()

    # Thêm keo Henkel và các mã FEFO tìm được
    extra_fefo = {
        "BIW71005512": "HSD (Henkel Keo SLoc 200C)",
        "BIW71005612": "HSD (Henkel Keo SLoc 200C)",
        "BIW30009362": "HSD",
        "BIW30111041": "HSD",
        "BIW32032012AA": "HSD",
        "BIW32032017AA": "HSD",
        "BIW32032018AA": "HSD",
        "CON69000038": "HSD",
    }
    for k, v in extra_fefo.items():
        if k not in fefo_materials:
            fefo_materials[k] = v
    for m in fefo_extra_from_td:
        if m not in fefo_materials:
            fefo_materials[m] = "HSD (Theo dõi cấp NVL)"

    # Box catalog from V31
    box_catalog = {}
    for r in range(5, ws_dm.max_row + 1):
        m = ws_dm.cell(r, 5).value
        if m:
            mat_key = str(m).strip().upper()
            if mat_key not in box_catalog:
                box_catalog[mat_key] = {
                    "model": str(ws_dm.cell(r, 4).value or "").strip(),
                    "material": mat_key,
                    "box_type": str(ws_dm.cell(r, 6).value or "").strip(),
                    "packing": safe_int_packing(ws_dm.cell(r, 7).value) or 1,
                    "ncc": ncc_map.get(mat_key, ""),
                }

    # Merge materials found in THEO DOI ODER into box_catalog if missing
    for mat_key, meta in material_meta_from_td.items():
        if mat_key not in box_catalog:
            box_catalog[mat_key] = {
                "model": meta.get("model", ""),
                "material": mat_key,
                "box_type": "Thùng tiêu chuẩn" if meta.get("packing", 0) > 0 else "Hàng rời/Box",
                "packing": meta.get("packing", 0),
                "ncc": ncc_map.get(mat_key, ""),
            }
        else:
            if not box_catalog[mat_key].get("ncc") and mat_key in ncc_map:
                box_catalog[mat_key]["ncc"] = ncc_map[mat_key]
            if not box_catalog[mat_key].get("model") and meta.get("model"):
                box_catalog[mat_key]["model"] = meta["model"]

    # Bin patterns from V31 + Physical warehouse rules
    existing_patterns = []
    for r in range(5, ws_dm.max_row + 1):
        bp = ws_dm.cell(r, 9).value
        bz = ws_dm.cell(r, 10).value
        if bp:
            existing_patterns.append({
                "pattern": str(bp).strip().upper(),
                "zone": str(bz or "KHAC").strip()
            })

    # Detailed warehouse physical bin rules specified by user
    warehouse_bin_rules = [
        {"pattern": "APFL-BL*", "zone": "BIN BLOCK", "desc": "Bin hệ thống khoá / block", "ground_level": "N/A"},
        {"pattern": "APFL-HO*", "zone": "Hàng Hold", "desc": "Hàng chờ xử lý chất lượng", "ground_level": "N/A"},
        {"pattern": "APFL-HTK*", "zone": "Hàng tập kết", "desc": "Khu vực tập kết hàng hoá", "ground_level": "Ground"},
        {"pattern": "APFL-MA*", "zone": "Hàng mới về", "desc": "Hàng mới tiếp nhận", "ground_level": "Ground"},
        {"pattern": "APFL-OV*", "zone": "MA2", "desc": "Khu vực tràn / MA2", "ground_level": "Ground"},
        {"pattern": "APFL-P22*", "zone": "Rack RV5", "desc": "Rack RV5 xếp chồng lên nhau", "ground_level": "Stackable Rack"},
        {"pattern": "APFL-P23*", "zone": "Rack RV5", "desc": "Rack RV5 xếp chồng lên nhau", "ground_level": "Stackable Rack"},
        {"pattern": "APFL-P24*", "zone": "Rack Cam", "desc": "Rack Cam xếp chồng lên nhau", "ground_level": "Stackable Rack"},
        {"pattern": "APFL-P25*", "zone": "Thùng giấy nhỏ", "desc": "Khu vực thùng giấy nhỏ", "ground_level": "Ground/Shelf"},
        {"pattern": "APFL-P26*", "zone": "Sọt sắt", "desc": "Khu vực sọt sắt", "ground_level": "Ground/Stack"},
        {"pattern": "APFL-P*", "zone": "Hàng Box", "desc": "Hàng Box để dưới mặt đất", "ground_level": "Ground Level 1"},
        {"pattern": "APFL-ZF16-02*", "zone": "ZF16 Mặt 2 Rack Cam", "desc": "ZF16 mặt 2 là rack cam", "ground_level": "Rack"},
        {"pattern": "APFL-ZF21*", "zone": "ZF21 Rack EDC", "desc": "ZF21 là rack EDC gần giống RV5", "ground_level": "Stackable Rack"},
        {"pattern": "APFL-ZF*", "zone": "Thùng giấy - Kiện gỗ", "desc": "ZF thùng giấy kiện gỗ xếp lên nhau", "ground_level": "Floor Stack"},
        {"pattern": "APFL-TH*", "zone": "Thùng giấy - Kiện gỗ", "desc": "TH thùng giấy kiện gỗ xếp lên nhau", "ground_level": "Floor Stack"},
        {"pattern": "APFL-HR1-*-*-1", "zone": "High Rack Tầng 1 (Mặt đất)", "desc": "Kệ cao tầng 1 mặt đất", "ground_level": "Ground Level 1"},
        {"pattern": "APFL-HR1-*-*-2", "zone": "High Rack Tầng 2 (Tầng trên)", "desc": "Kệ cao tầng 2", "ground_level": "Upper Level 2"},
        {"pattern": "APFL-HR1-*-*-3", "zone": "High Rack Tầng 3 (Tầng trên)", "desc": "Kệ cao tầng 3", "ground_level": "Upper Level 3"},
        {"pattern": "APFL-HR1-*-*-4", "zone": "High Rack Tầng 4 (Tầng trên)", "desc": "Kệ cao tầng 4", "ground_level": "Upper Level 4"},
        {"pattern": "APFL-HR*", "zone": "Box giấy-Hàng thùng", "desc": "Kệ cao High Rack tổng quát", "ground_level": "Multi-level"},
        {"pattern": "APFL-CC*", "zone": "Sọt sắt-Box giấy", "desc": "Khu vực Sọt sắt - Box giấy", "ground_level": "Ground"},
    ]

    # Merge patterns (specialized rules first, then general patterns)
    seen_patterns = set()
    final_bin_patterns = []
    for r in warehouse_bin_rules:
        pat = r["pattern"]
        if pat not in seen_patterns:
            seen_patterns.add(pat)
            final_bin_patterns.append(r)
    for p in existing_patterns:
        pat = p["pattern"]
        if pat not in seen_patterns:
            seen_patterns.add(pat)
            final_bin_patterns.append({
                "pattern": pat,
                "zone": p["zone"],
                "desc": "Kế thừa từ Master V31",
                "ground_level": "Ground" if "-1" in pat or "P" in pat else "Standard"
            })

    # Box rules from V31 CONFIG
    box_rules = []
    for r in range(3, ws_cfg.max_row + 1):
        pmin = ws_cfg.cell(r, 7).value
        pmax = ws_cfg.cell(r, 8).value
        nbox = ws_cfg.cell(r, 9).value
        act = ws_cfg.cell(r, 10).value
        if pmin is not None and str(act).upper() == "TRUE":
            box_rules.append({
                "min": safe_int_packing(pmin),
                "max": safe_int_packing(pmax),
                "boxes": safe_int_packing(nbox)
            })

    # 3. Extract FG Catalog from 261008_Stock_FG_MA_BODY.xlsm
    fg_path = base_dir / "261008_Stock_FG_MA_BODY.xlsm"
    wb_fg = openpyxl.load_workbook(str(fg_path), data_only=True)
    ws_fg_data = wb_fg["DATA"]

    fg_catalog = []
    print(f"[BUILD_MASTER] Extracting FG items from {fg_path.name}...")
    for r in range(2, ws_fg_data.max_row + 1):
        mat_fg = ws_fg_data.cell(r, 2).value
        if not mat_fg:
            continue
        model_fg = ws_fg_data.cell(r, 3).value
        us_fg = ws_fg_data.cell(r, 4).value
        tram_fg = ws_fg_data.cell(r, 5).value
        shop_fg = ws_fg_data.cell(r, 6).value
        ma_son = ws_fg_data.cell(r, 7).value
        khu_vuc = ws_fg_data.cell(r, 8).value
        try:
            us_val = float(us_fg) if us_fg is not None else 1.0
        except Exception:
            us_val = 1.0

        fg_catalog.append({
            "material": str(mat_fg).strip().upper(),
            "model": str(model_fg or "").strip(),
            "us": us_val,
            "tram": str(tram_fg or "").strip(),
            "shop": str(shop_fg or "BODY").strip(),
            "ma_son": str(ma_son or "").strip(),
            "khu_vuc": str(khu_vuc or "").strip(),
        })

    print(f"[BUILD_MASTER] Loaded {len(fg_catalog)} FG part items.")

    # 4. Create MASTER_DATA_VINFAST.xlsx with clear separation
    print(f"[BUILD_MASTER] Creating unified workbook {output_excel}...")
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    header_fill_nvl = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # Navy Blue
    header_fill_fg = PatternFill(start_color="375623", end_color="375623", fill_type="solid")   # Dark Green
    header_fill_cfg = PatternFill(start_color="595959", end_color="595959", fill_type="solid")  # Dark Gray
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Sheet 1: NVL_DANH_MUC
    ws1 = wb_out.create_sheet("NVL_DANH_MUC")
    ws1.views.sheetView[0].showGridLines = True
    headers_ws1 = ["STT", "Mã NVL", "Nhà Cung Cấp (NCC)", "Model", "Loại Thùng", "Quy Cách (Packing)", "Ghi Chú"]
    ws1.append(headers_ws1)
    for col_num in range(1, len(headers_ws1) + 1):
        cell = ws1.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_nvl
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sorted_boxes = sorted(box_catalog.values(), key=lambda x: x["material"])
    for idx, item in enumerate(sorted_boxes, 1):
        row_vals = [
            idx,
            item["material"],
            item.get("ncc", ""),
            item.get("model", ""),
            item.get("box_type", ""),
            item.get("packing", 0),
            "FEFO / HSD" if item["material"] in fefo_materials else ""
        ]
        ws1.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws1.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            if c in (1, 6):
                cell.alignment = Alignment(horizontal="center")

    # Sheet 2: NVL_FEFO_HSD
    ws2 = wb_out.create_sheet("NVL_FEFO_HSD")
    ws2.views.sheetView[0].showGridLines = True
    headers_ws2 = ["STT", "Mã Vật Tư", "Mô Tả / Đặc Điểm", "Loại HSD", "SLoc Đặc Thù", "Ghi Chú Phân Bổ"]
    ws2.append(headers_ws2)
    for col_num in range(1, len(headers_ws2) + 1):
        cell = ws2.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_nvl
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, (m_fefo, desc) in enumerate(sorted(fefo_materials.items()), 1):
        is_henkel = "200C" in desc or "Henkel" in desc or "BIW71005" in m_fefo
        row_vals = [
            idx,
            m_fefo,
            desc,
            "FEFO (Hạn sử dụng)",
            "200C" if is_henkel else "1005",
            "Ưu tiên số 1 tuyệt đối khi ra đơn. Keo kiểm tra đúng SLoc 200C" if is_henkel else "Ưu tiên số 1 tuyệt đối khi ra đơn"
        ]
        ws2.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws2.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            if c in (1, 4, 5):
                cell.alignment = Alignment(horizontal="center")

    # Sheet 3: NVL_BIN_ZONE
    ws3 = wb_out.create_sheet("NVL_BIN_ZONE")
    ws3.views.sheetView[0].showGridLines = True
    headers_ws3 = ["STT", "Mẫu Vị Trí (Pattern)", "Khu Vực (Zone)", "Đặc Điểm Lưu Trữ & Xếp Chồng", "Tầng / Cốt Sàn", "Thứ Tự Ưu Tiên Phụ"]
    ws3.append(headers_ws3)
    for col_num in range(1, len(headers_ws3) + 1):
        cell = ws3.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_nvl
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, pat in enumerate(final_bin_patterns, 1):
        row_vals = [
            idx,
            pat["pattern"],
            pat["zone"],
            pat.get("desc", ""),
            pat.get("ground_level", "Standard"),
            "1 - Đất / Dễ lấy" if "Ground" in pat.get("ground_level", "") or "Floor" in pat.get("ground_level", "") else "2 - Kệ cao / Cần xe nâng"
        ]
        ws3.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws3.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            if c in (1, 2, 5, 6):
                cell.alignment = Alignment(horizontal="center")

    # Sheet 4: NVL_CONFIG
    ws4 = wb_out.create_sheet("NVL_CONFIG")
    ws4.views.sheetView[0].showGridLines = True
    headers_ws4 = ["STT", "Min Packing", "Max Packing", "Số Thùng Tối Đa (V31)", "Trạng Thái"]
    ws4.append(headers_ws4)
    for col_num in range(1, len(headers_ws4) + 1):
        cell = ws4.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_cfg
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, rule in enumerate(box_rules, 1):
        row_vals = [idx, rule["min"], rule["max"], rule["boxes"], "TRUE"]
        ws4.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws4.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

    # Sheet 5: FG_DANH_MUC
    ws5 = wb_out.create_sheet("FG_DANH_MUC")
    ws5.views.sheetView[0].showGridLines = True
    headers_ws5 = ["STT", "Mã Chi Tiết (FG)", "Model", "US (Định Mức Xe)", "Trạm", "Xưởng (Shop)", "Mã Sau Sơn", "Khu Vực"]
    ws5.append(headers_ws5)
    for col_num in range(1, len(headers_ws5) + 1):
        cell = ws5.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_fg
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, item in enumerate(fg_catalog, 1):
        row_vals = [
            idx,
            item["material"],
            item["model"],
            item["us"],
            item["tram"],
            item["shop"],
            item["ma_son"],
            item["khu_vuc"]
        ]
        ws5.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws5.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            if c in (1, 4):
                cell.alignment = Alignment(horizontal="center")

    # Sheet 6: FG_SHOP_LINE
    ws6 = wb_out.create_sheet("FG_SHOP_LINE")
    ws6.views.sheetView[0].showGridLines = True
    headers_ws6 = ["STT", "Xưởng (Shop)", "Plant Kho", "Plant Line", "SLoc Kho (MA)", "SLoc Line", "Ngưỡng DOH Ca Chuẩn", "Ghi Chú"]
    ws6.append(headers_ws6)
    for col_num in range(1, len(headers_ws6) + 1):
        cell = ws6.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_fg
        cell.alignment = Alignment(horizontal="center", vertical="center")

    fg_shop_configs = [
        {"shop": "BODY", "plant_kho": "3093", "plant_line": "3093", "sloc_kho": "3093", "sloc_line": "3000", "doh": 200, "note": "Xuất bán thành phẩm sang Line Hàn"},
        {"shop": "EDC", "plant_kho": "3093", "plant_line": "3093", "sloc_kho": "3093", "sloc_line": "3000", "doh": 200, "note": "Xuất linh kiện xưởng EDC"},
        {"shop": "XMD", "plant_kho": "3093", "plant_line": "3093", "sloc_kho": "3093", "sloc_line": "3000", "doh": 200, "note": "Xuất xe máy điện / phụ tùng"},
    ]
    for idx, item in enumerate(fg_shop_configs, 1):
        row_vals = [idx, item["shop"], item["plant_kho"], item["plant_line"], item["sloc_kho"], item["sloc_line"], item["doh"], item["note"]]
        ws6.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws6.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            if c in (1, 2, 3, 4, 5, 6, 7):
                cell.alignment = Alignment(horizontal="center")

    # Sheet 7: FG_BIN_MAP
    ws7 = wb_out.create_sheet("FG_BIN_MAP")
    ws7.views.sheetView[0].showGridLines = True
    headers_ws7 = ["STT", "Mẫu Bin FG (Pattern)", "Loại Hàng Bán Thành Phẩm", "Quy Trình Điều Chuyển", "Thông Báo Sau Khi Tạo OB"]
    ws7.append(headers_ws7)
    for col_num in range(1, len(headers_ws7) + 1):
        cell = ws7.cell(1, col_num)
        cell.font = header_font
        cell.fill = header_fill_fg
        cell.alignment = Alignment(horizontal="center", vertical="center")

    fg_bin_rules = [
        {"pattern": "APFG-MA*", "type": "Hàng bán thành phẩm xuất xưởng khác", "flow": "Plant 3093 -> SLoc Line 3000", "msg": "Tự động gửi thông báo Teams sau khi tạo OB thành công"},
        {"pattern": "APFG-BODY*", "type": "Bán thành phẩm Body hàn", "flow": "Plant 3093 -> SLoc Line 3000", "msg": "Tự động gửi thông báo Teams sau khi hoàn tất"},
    ]
    for idx, item in enumerate(fg_bin_rules, 1):
        row_vals = [idx, item["pattern"], item["type"], item["flow"], item["msg"]]
        ws7.append(row_vals)
        for c in range(1, len(row_vals) + 1):
            cell = ws7.cell(idx + 1, c)
            cell.font = data_font
            cell.border = thin_border
            if c in (1, 2, 4):
                cell.alignment = Alignment(horizontal="center")

    # Auto-fit column widths
    for sheet in [ws1, ws2, ws3, ws4, ws5, ws6, ws7]:
        for col in sheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # Save Excel
    wb_out.save(str(output_excel))
    print(f"[BUILD_MASTER] Saved MASTER_DATA_VINFAST.xlsx ({os.path.getsize(str(output_excel))} bytes)")

    # 5. Build rich JSON cache for high speed O(1) in-memory lookup
    master_json_payload = {
        "version": "UNIFIED_MASTER_V20261027",
        "source_excel": "MASTER_DATA_VINFAST.xlsx",
        "fefo_materials": fefo_materials,
        "box_catalog": box_catalog,
        "bin_patterns": final_bin_patterns,
        "box_rules": box_rules,
        "vendor_map": ncc_map,
        "fg_catalog": {item["material"]: item for item in fg_catalog},
        "fg_shop_configs": {item["shop"]: item for item in fg_shop_configs},
    }

    with open(str(output_json), "w", encoding="utf-8") as f:
        json.dump(master_json_payload, f, ensure_ascii=False, indent=2)

    print(f"[BUILD_MASTER] Saved unified JSON cache to {output_json} ({os.path.getsize(str(output_json))} bytes)")
    print("[BUILD_MASTER] SUCCESS!")


if __name__ == "__main__":
    build_master_data()
